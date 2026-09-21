
# -*- coding: utf-8 -*-
import os, re
from datetime import datetime, timedelta, time


import xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

DIAS = ["lunes","martes","miércoles","jueves","viernes","sábado","domingo"]
OBS = [
    "Vacaciones","Permiso con goce","Permiso sin goce","Falta",
    "Falta injustificada","Tiempo por tiempo","Incapacidad","Comisión de trabajo","DESCANSO",
    "DOMINGO LABORADO","DESCANSO POR DOMINGO LABORADO","Otro"
]

def parse_times(v):
    return re.findall(r"\d{2}:\d{2}", "" if v in ("", None) else str(v))

def get_value(sh, row, label):
    for c in range(sh.ncols):
        if str(sh.cell_value(row, c)).strip() == label:
            for off in (2, 1):
                if c + off < sh.ncols and sh.cell_value(row, c + off) not in ("", None):
                    return str(sh.cell_value(row, c + off)).strip()
    return ""

def read_report(path):
    book = xlrd.open_workbook(path, formatting_info=False)
    sh = book.sheet_by_name("Reporte de Asistencia")

    period = ""
    for r in range(min(5, sh.nrows)):
        for c in range(min(8, sh.ncols)):
            v = sh.cell_value(r, c)
            if isinstance(v, str) and "~" in v:
                period = v
                break
        if period:
            break

    m = re.search(r"(\d{4}-\d{2}-\d{2})\s*~\s*(\d{4}-\d{2}-\d{2})", period)
    if not m:
        raise ValueError("No se pudo leer el periodo del StandardReport.")

    start = datetime.strptime(m.group(1), "%Y-%m-%d").date()
    end = datetime.strptime(m.group(2), "%Y-%m-%d").date()
    dates = [start + timedelta(days=i) for i in range((end - start).days + 1)]

    employees = []
    for r in range(sh.nrows):
        if str(sh.cell_value(r, 0)).strip() == "ID:":
            eid = str(sh.cell_value(r, 2)).strip()
            if eid.endswith(".0"):
                eid = eid[:-2]
            # IMPORTANT: ONLY the row immediately after the ID belongs to this employee.
            check_row = r + 1 if r + 1 < sh.nrows else None
            employees.append((eid, get_value(sh, r, "Nombre:"), get_value(sh, r, "Departamento:"), check_row))

    # Ordenar por departamento alfabéticamente y, dentro de cada departamento, por empleado.
    employees.sort(
        key=lambda x: (
            (x[2] or "").strip().casefold(),
            (x[1] or "").strip().casefold(),
            (x[0] or "").strip().casefold(),
        )
    )
    return sh, start, end, dates, employees

def build(sh, start, end, dates, employees, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Control de asistencia"
    ws.sheet_view.showGridLines = False

    blue, dark, yellow, green, gray, white = "4472C4","404040","FFF2CC","E2F0D9","E7E6E6","FFFFFF"
    thin = Side(style="thin", color="BFBFBF")

    ws.merge_cells("A1:G1")
    ws["A1"] = "CONTROL DE ASISTENCIA"
    ws["A1"].font = Font(bold=True, size=16, color=white)
    ws["A1"].fill = PatternFill("solid", fgColor=blue)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Periodo: {start:%d/%m/%Y} al {end:%d/%m/%Y}"
    ws["A2"].font = Font(bold=True, size=11)

    headers = ["Fecha","Día","Entrada","Salida","Horas trabajadas","Saldo","Observaciones"]
    row = 4

    for eid, name, dept, check_row in employees:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        ws.cell(row, 1).value = f"ID: {eid}    |    EMPLEADO: {name}    |    DEPARTAMENTO: {dept}"
        ws.cell(row, 1).font = Font(bold=True, color=white)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=blue)
        row += 1

        for c, h in enumerate(headers, 1):
            cell = ws.cell(row, c, h)
            cell.font = Font(bold=True, color=white)
            cell.fill = PatternFill("solid", fgColor=dark)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row += 1
        first = row
        obs_rows = []

        for idx, d in enumerate(dates):
            r = row
            ts = parse_times(sh.cell_value(check_row, idx)) if check_row is not None and idx < sh.ncols else []

            ws.cell(r, 1, d).number_format = "dd/mm/yyyy"
            ws.cell(r, 2, DIAS[d.weekday()])

            if ts:
                hh, mm = map(int, ts[0].split(":"))
                ws.cell(r, 3, time(hh, mm)).number_format = "hh:mm"
                if len(ts) >= 2:
                    hh, mm = map(int, ts[-1].split(":"))
                    ws.cell(r, 4, time(hh, mm)).number_format = "hh:mm"

            # H = hidden numeric saldo. This is what formulas use.
            ws.cell(r, 8).number_format = "0.000000"

            ws.cell(r, 5, f'=IF(OR(C{r}="",D{r}=""),0,MOD(D{r}-C{r},1))')
            ws.cell(r, 5).number_format = "[h]:mm"

            if d.weekday() == 6:
                ws.cell(r, 8, 0)
                ws.cell(r, 7, "DOMINGO LABORADO" if ts else "DESCANSO")
                if not ts:
                    ws.cell(r, 7).fill = PatternFill("solid", fgColor=gray)
            else:
                ws.cell(r, 8, (
                    f'=IF(OR(G{r}="Falta injustificada",G{r}="Tiempo por tiempo"),-TIME(8,0,0),'
                    f'IF(OR(G{r}="Vacaciones",G{r}="Permiso con goce",G{r}="Permiso sin goce",'
                    f'G{r}="Falta",G{r}="Incapacidad",G{r}="Comisión de trabajo"),0,'
                    f'IF(OR(C{r}="",D{r}=""),0,E{r}-TIME(8,0,0))))'
                ))
                ws.cell(r, 7).value = None

            # Visible signed saldo text, no negative-time formatting, no ####.
            ws.cell(r, 6, (
                f'=IF(H{r}=0,"0:00",IF(H{r}>0,"+"&TEXT(H{r},"[h]:mm"),'
                f'"-"&TEXT(ABS(H{r}),"[h]:mm")))'
            ))
            ws.cell(r, 6).number_format = "@"

            ws.cell(r, 3).fill = PatternFill("solid", fgColor=yellow)
            ws.cell(r, 4).fill = PatternFill("solid", fgColor=yellow)
            if not (d.weekday() == 6 and not ts):
                ws.cell(r, 7).fill = PatternFill("solid", fgColor=yellow)

            obs_rows.append(r)
            row += 1

        total = row
        ws.merge_cells(start_row=total, start_column=1, end_row=total, end_column=3)
        ws.cell(total, 1, "TOTAL HORAS TRABAJADAS:")
        ws.cell(total, 1).font = Font(bold=True)
        ws.cell(total, 4, f"=SUM(E{first}:E{row-1})")
        ws.cell(total, 4).number_format = "[h]:mm"
        ws.cell(total, 4).font = Font(bold=True)

        ws.cell(total, 5, "SALDO DEL PERIODO:")
        ws.cell(total, 5).font = Font(bold=True)

        ws.cell(total, 8, f"=SUM(H{first}:H{row-1})")
        ws.cell(total, 8).number_format = "0.000000"
        ws.cell(total, 6, (
            f'=IF(H{total}=0,"0:00",IF(H{total}>0,"+"&TEXT(H{total},"[h]:mm"),'
            f'"-"&TEXT(ABS(H{total}),"[h]:mm")))'
        ))
        ws.cell(total, 6).number_format = "@"
        ws.cell(total, 6).font = Font(bold=True)

        for c in range(1, 9):
            ws.cell(total, c).fill = PatternFill("solid", fgColor=green)

        row += 2

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row, 1, "Firma del empleado: __________________________________").font = Font(bold=True)
        ws.merge_cells(start_row=row, start_column=5, end_row=row, end_column=7)
        ws.cell(row, 5, "Firma del supervisor: ________________________________").font = Font(bold=True)
        row += 3

        dv = DataValidation(type="list", formula1='"' + ",".join(OBS) + '"', allow_blank=True)
        dv.error = "Selecciona una opción de la lista."
        dv.errorTitle = "Observación no válida"
        ws.add_data_validation(dv)
        for rr in obs_rows:
            dv.add(ws.cell(rr, 7))

    for rr in ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=8):
        for cell in rr:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    for i, w in enumerate([14,14,12,12,19,16,34,12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["H"].hidden = True

    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:2"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(out)

def main():
    import sys

    if len(sys.argv) < 3:
        print("ERROR: Faltan argumentos.")
        print("Uso: python generar_reporte.py archivo_entrada archivo_salida")
        sys.exit(1)

    src = sys.argv[1]
    out = sys.argv[2]

    try:
        sh, start, end, dates, employees = read_report(src)

        build(
            sh,
            start,
            end,
            dates,
            employees,
            out
        )

        print(
            f"OK|{start:%d/%m/%Y}|"
            f"{end:%d/%m/%Y}|"
            f"{len(employees)}"
        )

    except Exception as e:
        print(f"ERROR|{str(e)}")
        sys.exit(1)

if __name__ == "__main__":
    main()
